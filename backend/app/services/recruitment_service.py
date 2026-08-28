from __future__ import annotations

import csv
import io
import uuid
from collections.abc import Sequence
from typing import Any

from openpyxl import Workbook
from sqlalchemy import func, select

from app.core.exceptions import ConflictError, NotFoundError
from app.models.recruitment import (
    Candidate,
    CandidateDocument,
    CandidateSkill,
    CandidateSource,
    CandidateStageHistory,
    CandidateTalentPool,
    JobOpening,
    RecruiterNote,
    RecruitmentStage,
    TalentPool,
)
from app.models.requisition import JobRequisition, Notification
from app.repositories.recruitment_repository import (
    CandidateDocumentRepository,
    CandidateRepository,
    NoteRepository,
    OpeningRepository,
    PoolMemberRepository,
    SkillRepository,
    SourceRepository,
    StageHistoryRepository,
    StageRepository,
    TalentPoolRepository,
)
from app.repositories.requisition_repository import NotificationRepository
from app.schemas.recruitment import (
    CandidateCreate,
    CandidateListParams,
    CandidateUpdate,
    DocumentLink,
    NoteCreate,
    OpeningCreate,
    RecruitmentDashboard,
    StageCreate,
    StageMove,
    TalentPoolCreate,
)
from app.services.audit_service import AuditService
from app.utils.datetime import utc_now


class RecruitmentService:
    def __init__(
        self,
        opening: OpeningRepository,
        candidates: CandidateRepository,
        sources: SourceRepository,
        stages: StageRepository,
        documents: CandidateDocumentRepository,
        skills: SkillRepository,
        history: StageHistoryRepository,
        notes: NoteRepository,
        pools: TalentPoolRepository,
        members: PoolMemberRepository,
        notifications: NotificationRepository,
        audit: AuditService,
    ) -> None:
        self.opening: OpeningRepository = opening
        self.candidates: CandidateRepository = candidates
        self.sources: SourceRepository = sources
        self.stages: StageRepository = stages
        self.documents: CandidateDocumentRepository = documents
        self.skills: SkillRepository = skills
        self.history: StageHistoryRepository = history
        self.notes: NoteRepository = notes
        self.pools: TalentPoolRepository = pools
        self.members: PoolMemberRepository = members
        self.notifications: NotificationRepository = notifications
        self.audit: AuditService = audit

    async def create_opening(self, p: OpeningCreate, actor: uuid.UUID) -> JobOpening:
        req = await self.opening.session.get(JobRequisition, p.requisition_id)
        if not req or req.deleted_at:
            raise NotFoundError("Job requisition")
        if req.status != "approved":
            raise ConflictError("Only approved requisitions can become job openings.")
        if await self.opening.get_by(requisition_id=p.requisition_id):
            raise ConflictError("A job opening already exists for this requisition.")
        row = await self.opening.add(JobOpening(**p.model_dump()), actor_id=actor)
        await self._audit("opening.created", row.id, actor)
        return row

    async def opening_action(self, opening_id: uuid.UUID, action: str, actor: uuid.UUID) -> JobOpening:
        row = await self.opening.get(opening_id)
        if not row:
            raise NotFoundError("Job opening")
        transitions = {
            "publish": ({"draft"}, "published"),
            "close": ({"published"}, "closed"),
            "cancel": ({"draft", "published"}, "cancelled"),
            "fill": ({"published"}, "filled"),
        }
        allowed, target = transitions[action]
        if row.status not in allowed:
            raise ConflictError(f"A {row.status} opening cannot be {action}ed.")
        values: dict[str, Any] = {"status": target}
        if target == "published":
            values["published_at"] = utc_now()
        if target in {"closed", "filled"}:
            values["closed_at"] = utc_now()
        await self.opening.update(row, values, actor_id=actor)
        if target == "closed" and row.recruiter_id:
            await self._notify(
                row.recruiter_id, "Job opening closed", row.job_code, f"/recruitment/openings/{row.id}", actor
            )
        await self._audit(f"opening.{action}", row.id, actor)
        return row

    async def list_openings(self) -> Sequence[JobOpening]:
        return await self.opening.all()

    async def create_candidate(self, p: CandidateCreate, actor: uuid.UUID) -> Candidate:
        opening = await self.opening.get(p.job_opening_id)
        if not opening or opening.status != "published":
            raise ConflictError("Candidates require a published job opening.")
        duplicates = await self.candidates.duplicates(p.email, p.mobile_number)
        if duplicates:
            raise ConflictError(
                "A candidate with this email or mobile already exists.", error_code="duplicate_candidate"
            )
        if not await self.sources.get(p.source_id):
            raise NotFoundError("Candidate source")
        stage = await self.stages.get_by(name="Applied")
        if not stage:
            raise NotFoundError("Applied recruitment stage")
        if not await self.documents.document_exists(p.resume_document_id):
            raise NotFoundError("Resume document")
        values = p.model_dump(exclude={"skills", "resume_document_id"})
        if values.get("linkedin_url"):
            values["linkedin_url"] = str(values["linkedin_url"])
        row = await self.candidates.add(Candidate(**values, stage_id=stage.id), actor_id=actor)
        for name in dict.fromkeys(x.strip() for x in p.skills if x.strip()):
            await self.skills.add(CandidateSkill(candidate_id=row.id, name=name), actor_id=actor)
        await self.documents.add(
            CandidateDocument(candidate_id=row.id, document_id=p.resume_document_id, document_kind="resume"),
            actor_id=actor,
        )
        await self.history.add(
            CandidateStageHistory(candidate_id=row.id, to_stage_id=stage.id, comments="Candidate applied"),
            actor_id=actor,
        )
        await self._notify(
            opening.recruiter_id or actor,
            "New candidate added",
            f"{row.candidate_code}: {row.first_name} {row.last_name}",
            f"/recruitment/candidates/{row.id}",
            actor,
        )
        await self._audit("candidate.created", row.id, actor)
        return await self.get_candidate(row.id)

    async def get_candidate(self, candidate_id: uuid.UUID) -> Candidate:
        row = await self.candidates.get_detailed(candidate_id)
        if not row:
            raise NotFoundError("Candidate")
        return row

    async def update_candidate(
        self, candidate_id: uuid.UUID, p: CandidateUpdate, actor: uuid.UUID
    ) -> Candidate:
        row = await self.get_candidate(candidate_id)
        values = p.model_dump(exclude_unset=True)
        if values.get("linkedin_url"):
            values["linkedin_url"] = str(values["linkedin_url"])
        await self.candidates.update(row, values, actor_id=actor)
        await self._audit("candidate.updated", row.id, actor)
        return row

    async def search(self, p: CandidateListParams) -> tuple[Sequence[Candidate], int]:
        return await self.candidates.search(p)

    async def duplicates(self, email: str, mobile: str) -> Sequence[Candidate]:
        return await self.candidates.duplicates(email, mobile)

    async def move_stage(self, candidate_id: uuid.UUID, p: StageMove, actor: uuid.UUID) -> Candidate:
        row = await self.get_candidate(candidate_id)
        stage = await self.stages.get(p.stage_id)
        if not stage or not stage.is_active:
            raise NotFoundError("Recruitment stage")
        old = row.stage_id
        if old == stage.id:
            raise ConflictError("Candidate is already in this stage.")
        await self.candidates.update(
            row,
            {"stage_id": stage.id, "hired_at": utc_now() if stage.name == "Joined" else row.hired_at},
            actor_id=actor,
        )
        await self.history.add(
            CandidateStageHistory(
                candidate_id=row.id, from_stage_id=old, to_stage_id=stage.id, comments=p.comments
            ),
            actor_id=actor,
        )
        await self._notify(
            row.recruiter_id or actor,
            "Candidate stage changed",
            f"{row.candidate_code} moved to {stage.name}",
            f"/recruitment/candidates/{row.id}",
            actor,
        )
        await self._audit("candidate.stage_moved", row.id, actor)
        return await self.get_candidate(row.id)

    async def link_document(
        self, candidate_id: uuid.UUID, p: DocumentLink, actor: uuid.UUID
    ) -> CandidateDocument:
        await self.get_candidate(candidate_id)
        if not await self.documents.document_exists(p.document_id):
            raise NotFoundError("Document")
        item = await self.documents.add(
            CandidateDocument(candidate_id=candidate_id, **p.model_dump()), actor_id=actor
        )
        await self._audit("candidate.document_uploaded", candidate_id, actor)
        return item

    async def add_note(self, candidate_id: uuid.UUID, p: NoteCreate, actor: uuid.UUID) -> RecruiterNote:
        await self.get_candidate(candidate_id)
        item = await self.notes.add(
            RecruiterNote(candidate_id=candidate_id, **p.model_dump()), actor_id=actor
        )
        await self._audit("candidate.note_added", candidate_id, actor)
        return item

    async def references(self) -> tuple[Sequence[CandidateSource], Sequence[RecruitmentStage]]:
        return await self.sources.list(limit=100, order_by="name", descending=False), await self.stages.list(
            limit=100, order_by="sequence", descending=False
        )

    async def create_stage(self, p: StageCreate, actor: uuid.UUID) -> RecruitmentStage:
        if await self.stages.find(
            (func.lower(RecruitmentStage.name) == p.name.lower()) | (RecruitmentStage.sequence == p.sequence)
        ):
            raise ConflictError("Stage name or sequence already exists.")
        item = await self.stages.add(RecruitmentStage(**p.model_dump()), actor_id=actor)
        await self._audit("recruitment.stage_created", item.id, actor)
        return item

    async def list_pools(self) -> Sequence[TalentPool]:
        return await self.pools.list(limit=500, order_by="name", descending=False)

    async def create_pool(self, p: TalentPoolCreate, actor: uuid.UUID) -> TalentPool:
        if await self.pools.find(func.lower(TalentPool.name) == p.name.lower()):
            raise ConflictError("Talent pool name already exists.")
        return await self.pools.add(TalentPool(**p.model_dump()), actor_id=actor)

    async def add_pool_member(
        self, pool_id: uuid.UUID, candidate_id: uuid.UUID, actor: uuid.UUID
    ) -> CandidateTalentPool:
        if not await self.pools.get(pool_id):
            raise NotFoundError("Talent pool")
        await self.get_candidate(candidate_id)
        if await self.members.get_by(talent_pool_id=pool_id, candidate_id=candidate_id):
            raise ConflictError("Candidate is already in this talent pool.")
        return await self.members.add(
            CandidateTalentPool(talent_pool_id=pool_id, candidate_id=candidate_id), actor_id=actor
        )

    async def dashboard(self) -> RecruitmentDashboard:
        s = self.opening.session

        async def count(model: Any, *where: Any) -> int:
            return int(
                await s.scalar(
                    select(func.count()).select_from(model).where(model.deleted_at.is_(None), *where)
                )
                or 0
            )

        by_stage = (
            await s.execute(
                select(RecruitmentStage.name, func.count(Candidate.id))
                .outerjoin(
                    Candidate, (Candidate.stage_id == RecruitmentStage.id) & Candidate.deleted_at.is_(None)
                )
                .where(RecruitmentStage.deleted_at.is_(None))
                .group_by(RecruitmentStage.id, RecruitmentStage.name)
                .order_by(RecruitmentStage.sequence)
            )
        ).all()
        workload = (
            await s.execute(
                select(Candidate.recruiter_id, func.count())
                .where(Candidate.deleted_at.is_(None))
                .group_by(Candidate.recruiter_id)
            )
        ).all()
        offer_released = await self.stages.get_by(name="Offer Released")
        offer_accepted = await self.stages.get_by(name="Offer Accepted")
        joining = await self.stages.get_by(name="Joining Pending")
        rejected = await self.stages.get_by(name="Rejected")
        trend_rows = (
            await s.execute(
                select(
                    func.extract("year", Candidate.applied_at),
                    func.extract("month", Candidate.applied_at),
                    func.count(),
                )
                .where(Candidate.deleted_at.is_(None))
                .group_by(
                    func.extract("year", Candidate.applied_at),
                    func.extract("month", Candidate.applied_at),
                )
                .order_by(
                    func.extract("year", Candidate.applied_at),
                    func.extract("month", Candidate.applied_at),
                )
            )
        ).all()
        time_to_hire = await s.scalar(
            select(func.avg(func.extract("epoch", Candidate.hired_at - Candidate.applied_at) / 86400)).where(
                Candidate.deleted_at.is_(None), Candidate.hired_at.is_not(None)
            )
        )
        time_to_fill = await s.scalar(
            select(
                func.avg(func.extract("epoch", JobOpening.closed_at - JobOpening.published_at) / 86400)
            ).where(
                JobOpening.deleted_at.is_(None),
                JobOpening.closed_at.is_not(None),
                JobOpening.published_at.is_not(None),
            )
        )
        return RecruitmentDashboard(
            open_jobs=await count(JobOpening, JobOpening.status == "published"),
            total_applicants=await count(Candidate),
            offers_released=(
                await count(Candidate, Candidate.stage_id == offer_released.id) if offer_released else 0
            ),
            offers_accepted=(
                await count(Candidate, Candidate.stage_id == offer_accepted.id) if offer_accepted else 0
            ),
            offers_declined=await count(Candidate, Candidate.stage_id == rejected.id) if rejected else 0,
            joining_pending=await count(Candidate, Candidate.stage_id == joining.id) if joining else 0,
            by_stage=[{"label": n, "count": int(c)} for n, c in by_stage],
            recruiter_workload=[{"label": str(u or "Unassigned"), "count": int(c)} for u, c in workload],
            hiring_trend=[
                {"label": f"{int(year):04d}-{int(month):02d}", "count": int(total)}
                for year, month, total in trend_rows
            ],
            average_time_to_hire_days=round(float(time_to_hire or 0), 1),
            average_time_to_fill_days=round(float(time_to_fill or 0), 1),
        )

    async def export(self, p: CandidateListParams, fmt: str) -> tuple[bytes, str]:
        rows, _ = await self.candidates.search(p.model_copy(update={"page": 1, "page_size": 100}))
        headers = [
            "Candidate ID",
            "Name",
            "Email",
            "Mobile",
            "Stage",
            "Source",
            "Experience",
            "Notice Period",
        ]
        data = [
            [
                r.candidate_code,
                f"{r.first_name} {r.last_name}",
                r.email,
                r.mobile_number,
                r.stage.name,
                r.source.name,
                r.experience_years,
                r.notice_period_days,
            ]
            for r in rows
        ]
        if fmt == "csv":
            text = io.StringIO()
            writer = csv.writer(text)
            writer.writerow(headers)
            writer.writerows(data)
            return text.getvalue().encode(), "text/csv"
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for item in data:
            ws.append(item)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    async def _notify(
        self,
        user: uuid.UUID | None,
        title: str,
        message: str,
        link: str,
        actor: uuid.UUID,
    ) -> None:
        if user:
            await self.notifications.add(
                Notification(
                    user_id=user, title=title, message=message, link=link, notification_type="recruitment"
                ),
                actor_id=actor,
            )

    async def _audit(self, action: str, entity_id: uuid.UUID, actor: uuid.UUID) -> None:
        await self.audit.record_success(
            action, actor_id=actor, entity_type="recruitment", entity_id=entity_id
        )
